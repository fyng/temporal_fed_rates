"""Tests for fedrates.sources: cache behaviour, failure modes and parsers."""

from __future__ import annotations

import os
import re
import time
from datetime import datetime
from pathlib import Path

import pandas as pd
import pytest
import requests

from fedrates.sources import HLW_URL, RAW_DIR, WU_XIA_URL, download, hlw, wu_xia

# --- Fixtures and helpers ----------------------------------------------------

_REAL_GET = requests.get  # captured at import, before any fixture can patch it


def _no_network(*args: object, **kwargs: object) -> None:
    raise requests.ConnectionError("network disabled in tests")


@pytest.fixture
def no_network(monkeypatch: pytest.MonkeyPatch) -> None:
    """Block requests.get; shadows any same-named conftest fixture."""
    monkeypatch.setattr(requests, "get", _no_network)


@pytest.fixture
def network(monkeypatch: pytest.MonkeyPatch) -> None:
    """Restore real requests.get, undoing conftest's autouse network block."""
    monkeypatch.setattr(requests, "get", _REAL_GET)


class _Response:
    def __init__(self, content: bytes) -> None:
        self.content = content

    def raise_for_status(self) -> None:
        return None


def _age_file(path: Path, days: float) -> None:
    mtime = time.time() - days * 86400
    os.utime(path, (mtime, mtime))


def _write_wu_xia(path: Path) -> None:
    """Build a minimal Wu-Xia xlsx matching the observed Atlanta Fed layout."""
    import openpyxl  # optional at test time until it lands in pyproject

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append([None, "Effective federal funds rate (%", "Wu-Xia shadow federal funds rate"])
    for row in [
        (datetime(1960, 1, 1), 4.0, None),
        (datetime(1960, 2, 1), 4.0, None),
        (datetime(1990, 1, 1), 8.25, 7.95),
        (datetime(1990, 2, 1), 8.5, 7.9),
    ]:
        ws.append(row)
    wb.save(path)


def _write_hlw(path: Path) -> None:
    """Build a minimal HLW xlsx matching the observed New York Fed layout."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HLW Estimates"
    ws["A1"] = "This spreadsheet contains updated estimates"
    ws["C5"] = "Trend Growth (g), Annualized"
    ws["G5"] = "Other Determinants (z)"
    ws["K5"] = "Natural Rate (r*)"
    ws.append(
        ["Date", None, "US", "Canada", "Euro Area", None, "US", "Canada", "Euro Area",
         None, "US", "Canada", "Euro Area"]
    )
    for date, g, z, r_star in [
        (datetime(1961, 1, 1), 4.92, 0.005, 5.48),
        (datetime(1961, 4, 1), 5.11, 0.011, 5.70),
        (datetime(1961, 7, 1), 5.03, 0.013, 5.90),
    ]:
        ws.append([date, None, g, 3.17, "NA", None, z, -0.04, "NA", None, r_star, 3.55, "NA"])
    wb.save(path)


# --- download ----------------------------------------------------------------


def test_raw_dir_is_repo_data_raw() -> None:
    assert RAW_DIR == Path(__file__).resolve().parents[1] / "data" / "raw"


def test_download_serves_fresh_file_without_fetch(tmp_path: Path, no_network: None) -> None:
    path = tmp_path / "f.xlsx"
    path.write_bytes(b"cached")
    assert download("https://example.com/f.xlsx", path) == path
    assert path.read_bytes() == b"cached"


def test_download_fresh_within_max_age_skips_fetch(tmp_path: Path, no_network: None) -> None:
    path = tmp_path / "f.xlsx"
    path.write_bytes(b"cached")
    _age_file(path, 29.0)
    download("https://example.com/f.xlsx", path, max_age_days=30.0)
    assert path.read_bytes() == b"cached"


def test_download_stale_beyond_max_age_refetches(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "f.xlsx"
    path.write_bytes(b"stale")
    _age_file(path, 31.0)
    monkeypatch.setattr(requests, "get", lambda *a, **k: _Response(b"fresh"))
    assert download("https://example.com/f.xlsx", path, max_age_days=30.0).read_bytes() == b"fresh"


def test_download_force_refetches_fresh_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "f.xlsx"
    path.write_bytes(b"cached")
    monkeypatch.setattr(requests, "get", lambda *a, **k: _Response(b"fresh"))
    assert download("https://example.com/f.xlsx", path, force=True).read_bytes() == b"fresh"


def test_download_falls_back_to_cached_when_refresh_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "f.xlsx"
    path.write_bytes(b"cached")
    monkeypatch.setattr(requests, "get", _no_network)
    assert download("https://example.com/f.xlsx", path, max_age_days=0.0) == path
    assert path.read_bytes() == b"cached"


def test_download_missing_without_network_raises(tmp_path: Path, no_network: None) -> None:
    with pytest.raises(FileNotFoundError, match="example.com"):
        download("https://example.com/data/f.xlsx", tmp_path / "f.xlsx")


def test_download_http_error_without_cached_file_raises(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    class _R:
        content = b""

        def raise_for_status(self) -> None:
            raise requests.HTTPError("404")

    monkeypatch.setattr(requests, "get", lambda *a, **k: _R())
    with pytest.raises(FileNotFoundError, match="example.com"):
        download("https://example.com/data/f.xlsx", tmp_path / "f.xlsx")


# --- loader failure mode -----------------------------------------------------


def test_wu_xia_missing_file_raises_with_url(tmp_path: Path, no_network: None) -> None:
    with pytest.raises(FileNotFoundError, match=re.escape(WU_XIA_URL)):
        wu_xia(path=tmp_path / "WuXiaShadowRate.xlsx")


def test_hlw_missing_file_raises_with_url(tmp_path: Path, no_network: None) -> None:
    with pytest.raises(FileNotFoundError, match=re.escape(HLW_URL)):
        hlw(path=tmp_path / "hlw.xlsx")


# --- parsers -----------------------------------------------------------------


def test_wu_xia_parses_synthetic_layout(tmp_path: Path) -> None:
    path = tmp_path / "WuXiaShadowRate.xlsx"
    _write_wu_xia(path)
    s = wu_xia(path=path)
    assert s.name == "shadow_rate"
    assert s.index.name == "date"
    assert isinstance(s.index, pd.DatetimeIndex)
    assert s.dtype == "float64"
    assert list(s.index) == [
        pd.Timestamp("1960-01-01"),
        pd.Timestamp("1960-02-01"),
        pd.Timestamp("1990-01-01"),
        pd.Timestamp("1990-02-01"),
    ]
    assert s.iloc[:2].isna().all()
    assert s.iloc[2:].to_list() == pytest.approx([7.95, 7.9])


def test_hlw_parses_synthetic_layout(tmp_path: Path) -> None:
    path = tmp_path / "hlw.xlsx"
    _write_hlw(path)
    r = hlw(path=path)
    assert r.name == "r_star"
    assert r.index.name == "date"
    assert isinstance(r.index, pd.DatetimeIndex)
    assert r.dtype == "float64"
    assert list(r.index) == [
        pd.Timestamp("1961-01-01"),
        pd.Timestamp("1961-04-01"),
        pd.Timestamp("1961-07-01"),
    ]
    assert r.to_list() == pytest.approx([5.48, 5.7, 5.9])
    g = hlw(path=path, measure="g")
    assert g.name == "g"
    assert g.to_list() == pytest.approx([4.92, 5.11, 5.03])
    z = hlw(path=path, measure="z")
    assert z.name == "z"
    assert z.to_list() == pytest.approx([0.005, 0.011, 0.013])


def test_hlw_bad_measure_raises(tmp_path: Path) -> None:
    path = tmp_path / "hlw.xlsx"
    _write_hlw(path)
    with pytest.raises(KeyError, match="measure"):
        hlw(path=path, measure="nope")


# --- real files (network) ----------------------------------------------------


@pytest.mark.network
def test_network_wu_xia_real_file(tmp_path: Path, network: None) -> None:
    s = wu_xia(path=tmp_path / "WuXiaShadowRate.xlsx")
    assert s.index.inferred_freq == "MS"
    assert s.index.min() <= pd.Timestamp("1970-01-01")
    assert s.loc["2014"].min() < -2.0


@pytest.mark.network
def test_network_hlw_real_file(tmp_path: Path, network: None) -> None:
    r = hlw(path=tmp_path / "hlw.xlsx")
    assert set(r.index.month) == {1, 4, 7, 10}
    assert (r.index.day == 1).all()
    assert r.index.min() <= pd.Timestamp("1970-01-01")
    assert -1.5 < r.min() < r.max() < 6.0
